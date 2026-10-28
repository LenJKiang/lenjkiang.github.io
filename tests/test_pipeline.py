"""Synthetic fixtures only; never copied into production data or public output."""
import json, shutil, sys, tempfile, unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from common import ROOT, load, dump, sha, validate, base_path
from scan_resources import scan
from review_resource import decide
from generate_pages import generate
class ResourceContracts(unittest.TestCase):
 def setUp(self):
  self.temp=tempfile.TemporaryDirectory(prefix='test-case-',dir=ROOT/'workspace/working');self.r=Path(self.temp.name)
  for d in ['data','config','workspace/inbox/papers','workspace/inbox/teaching','workspace/review','public/papers','website/shared']:(self.r/d).mkdir(parents=True,exist_ok=True)
  for name in ['profile','research','publications','courses','resources','links']:shutil.copy2(ROOT/f'data/{name}.yml',self.r/f'data/{name}.yml')
  for name in ['site','public-assets']:shutil.copy2(ROOT/f'config/{name}.yml',self.r/f'config/{name}.yml')
  for f in (ROOT/'website/shared').glob('*'):shutil.copy2(f,self.r/'website/shared'/f.name)
  dump(self.r/'workspace/review/approvals.yml',[])
  dump(self.r/'data/resources.yml',[]);dump(self.r/'config/public-assets.yml',[])
  pubs=load(self.r/'data/publications.yml')
  for pub in pubs:pub.update(pdf_public_path=None,pdf_permission_status='external_link_only',pdf_version='unknown')
  pubs[0]['doi']='10.9999/synthetic-test'
  dump(self.r/'data/publications.yml',pubs)
  courses=load(self.r/'data/courses.yml')
  for course in courses:course['resources']=[]
  dump(self.r/'data/courses.yml',courses)
  profile=load(self.r/'data/profile.yml');profile['photo_resource_id']=None;dump(self.r/'data/profile.yml',profile)
  import fitz
  self.f=self.r/'workspace/inbox/papers/synthetic.pdf';self.doi=load(self.r/'data/publications.yml')[0]['doi']
  doc=fitz.open();p=doc.new_page();p.insert_text((72,72),'SYNTHETIC_PRIVATE_FIXTURE\nDOI reference for matching test: '+str(self.doi));doc.set_metadata({'title':'Synthetic fixture for pipeline testing'});doc.save(self.f);doc.close()
 def tearDown(self):self.temp.cleanup()
 def approve(self,**overrides):
  params=dict(root=self.r,source=self.f.relative_to(self.r).as_posix(),resource_id='synthetic-test',status='approved',expected_hash=sha(self.f),related=load(self.r/'data/publications.yml')[0]['id'],version='accepted_manuscript',language='en',license='synthetic-test-only',evidence='Synthetic test fixture, never a real release approval',approver='TEST FIXTURE',confirmed=True)
  params.update(overrides);return decide(**params)
 def test_scanner_preserves_bytes(self):
  before=sha(self.f);report=scan(self.r);self.assertEqual(sha(self.f),before);self.assertEqual(report['items'][0]['permission_status'],'pending_review');self.assertEqual(load(self.r/'data/resources.yml'),[])
 def test_pdf_doi_match(self):self.assertEqual(scan(self.r)['items'][0]['matched_publications'],[load(self.r/'data/publications.yml')[0]['id']])
 def test_exact_duplicate(self):
  shutil.copy2(self.f,self.f.with_name('copy.pdf'));self.assertEqual(len(scan(self.r)['exact_duplicates']),1)
 def test_metadata_duplicate_different_hashes(self):
  import fitz
  doc=fitz.open(self.f);doc[0].insert_text((72,130),'Different version');doc.save(self.f.with_name('version2.pdf'));doc.close();report=scan(self.r);self.assertEqual(len(report['possible_metadata_duplicates']),1);self.assertEqual(report['exact_duplicates'],[])
 def test_pptx_course_and_lecture(self):
  from pptx import Presentation
  p=self.r/'workspace/inbox/teaching/probability-lecture-1.pptx';prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[0]);slide.shapes.title.text='Probability Theory Lecture 1';prs.core_properties.title='Probability Theory Lecture 1';prs.save(p)
  item=scan(self.r,'teaching')['items'][0];self.assertIn('probability',item['possible_courses']);self.assertEqual(item['page_or_slide_count'],1);self.assertIsNotNone(item['chapter_candidate'])
 def test_teaching_pptx_publication_rejected(self):
  from pptx import Presentation
  f=self.r/'workspace/inbox/teaching/synthetic.pptx';Presentation().save(f)
  with self.assertRaises(AssertionError):self.approve(source=f.relative_to(self.r).as_posix(),expected_hash=sha(f),related='probability')
  self.assertFalse(list((self.r/'public').rglob('*.pptx')))
 def test_corrupt_pdf_graceful(self):
  (self.f.parent/'broken.pdf').write_bytes(b'not a pdf');items=scan(self.r)['items'];self.assertTrue(next(x for x in items if x['source_file'].endswith('broken.pdf'))['errors'])
 def test_no_confirmation_rejected(self):
  with self.assertRaises(AssertionError):self.approve(confirmed=False)
  self.assertEqual(list((self.r/'public').rglob('*.pdf')),[])
 def test_missing_license_rejected(self):
  with self.assertRaises(AssertionError):self.approve(license=None)
 def test_hash_change_rejected(self):
  with self.assertRaises(AssertionError):self.approve(expected_hash='0'*64)
 def test_unknown_paper_version_rejected(self):
  with self.assertRaises(AssertionError):self.approve(version='unknown')
 def test_private_decision_never_copies(self):
  self.approve(status='private_only');self.assertFalse(list((self.r/'public').rglob('*.pdf')));self.assertEqual(load(self.r/'data/resources.yml'),[])
 def test_approved_exact_copy_and_links(self):
  self.approve();d=validate(self.r);a=d['resources'][0];self.assertEqual(sha(self.r/a['public_path']),sha(self.f));generate(d,load(self.r/'config/site.yml'),self.r)
  for lang in ['zh','en']:
   s=(self.r/f'website/{lang}/publications'/f"{d['publications'][0]['id']}.qmd").read_text(encoding='utf-8');self.assertIn(' download ',s);self.assertIn('application/pdf',s)
 def test_public_tamper_fails(self):
  self.approve();a=load(self.r/'data/resources.yml')[0];(self.r/a['public_path']).write_bytes(b'changed')
  with self.assertRaises(AssertionError):validate(self.r)
 def test_source_change_invalidates(self):
  self.approve();self.f.write_bytes(b'changed original')
  with self.assertRaises(AssertionError):validate(self.r)
 def test_unapproved_binary_blocks_build(self):
  (self.r/'public/papers/rogue.pdf').write_bytes(b'private')
  with self.assertRaises(AssertionError):validate(self.r)
 def test_duplicate_publication_rejected(self):
  pubs=load(self.r/'data/publications.yml');p=dict(pubs[0]);p['id']='duplicate-id';pubs.append(p);dump(self.r/'data/publications.yml',pubs)
  with self.assertRaises(AssertionError):validate(self.r)
 def test_broken_resource_reference_rejected(self):
  courses=load(self.r/'data/courses.yml');courses[0]['resources']=['missing'];dump(self.r/'data/courses.yml',courses)
  with self.assertRaises(AssertionError):validate(self.r)
 def test_invalid_year_type_rejected(self):
  pubs=load(self.r/'data/publications.yml');pubs[0]['year']='2025';dump(self.r/'data/publications.yml',pubs)
  with self.assertRaises(AssertionError):validate(self.r)
 def test_private_path_in_metadata_rejected(self):
  p=load(self.r/'data/profile.yml');p['private_note']='E:/private/example.pdf';dump(self.r/'data/profile.yml',p)
  with self.assertRaises(AssertionError):validate(self.r)
 def test_path_traversal_rejected(self):
  with self.assertRaises(AssertionError):self.approve(source='../../outside.pdf')
 def test_both_base_path_formats(self):
  self.assertEqual(base_path('/'),'/');self.assertEqual(base_path('/Academic_Website/'),'/Academic_Website/')
  for bad in ['repo','/repo','//','/../','/repo/?x']:
   with self.assertRaises(AssertionError):base_path(bad)
 def test_unsigned_manifest_rejected(self):
  self.approve();dump(self.r/'config/public-assets.yml',[])
  with self.assertRaises(AssertionError):validate(self.r)
if __name__=='__main__':unittest.main(verbosity=2)
